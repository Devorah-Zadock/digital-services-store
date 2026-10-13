"""Shared building blocks for DeskKit's downloadable Excel templates.

Every template is written once, with each piece of text given in Hebrew and
English (T(he, en)); build_xlsx.py renders both languages. Conventions the
site's product pages promise:
  * yellow cells = the only cells to fill in; everything else is formulas
  * sheets are right-to-left in Hebrew, left-to-right in English
  * plenty of ready rows (formulas already in place), totals ABOVE the table
    so the table can grow without breaking them
  * dropdowns read from an editable "Lists" sheet
"""
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as col
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.formatting.rule import CellIsRule, FormulaRule, DataBarRule

GREEN = "1F5C4E"
GREEN_LIGHT = "E8F1EE"
INPUT = "FFF9C4"
INPUT_FONT = "0000CC"
GOLD = "C9973A"
GREY = "666666"
INK = "222222"
RED_FILL = "FDE2E1"
RED_FONT = "B42318"
OK_FILL = "DCF2E3"
OK_FONT = "1E6B3A"
AMBER_FILL = "FFF1D6"
AMBER_FONT = "8A5A00"
LINE = Side(style="thin", color="D9DED9")
BOX = Border(left=LINE, right=LINE, top=LINE, bottom=LINE)
FONT = "Arial"


class Lang:
    def __init__(self, code):
        self.code = code
        self.he = code == "he"

    def __call__(self, he, en):
        return he if self.he else en

    @property
    def money(self):
        return '#,##0 "₪";[Red]-#,##0 "₪"' if self.he else '#,##0;[Red]-#,##0'

    @property
    def money2(self):
        return '#,##0.00 "₪";[Red]-#,##0.00 "₪"' if self.he else '#,##0.00;[Red]-#,##0.00'

    @property
    def cur(self):
        """Currency marker for column headings."""
        return " (₪)" if self.he else ""

    date = "dd/mm/yyyy"
    time = "hh:mm"
    pct = "0%"
    pct1 = "0.0%"
    num = "#,##0"
    num1 = "#,##0.0"
    num2 = "#,##0.00"


class Book:
    def __init__(self, lang):
        self.L = lang
        self.wb = Workbook()
        self.wb.remove(self.wb.active)
        self.wb.calculation.fullCalcOnLoad = True

    def sheet(self, title, widths, landscape=True):
        ws = self.wb.create_sheet(title[:31])
        ws.sheet_view.rightToLeft = self.L.he
        ws.sheet_view.showGridLines = False
        ws.sheet_view.zoomScale = 100
        for i, w in enumerate(widths, start=1):
            ws.column_dimensions[col(i)].width = w
        ws.page_setup.orientation = "landscape" if landscape else "portrait"
        ws.page_setup.paperSize = ws.PAPERSIZE_A4
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.print_options.horizontalCentered = True
        ws.page_margins.left = ws.page_margins.right = 0.4
        return Sheet(self, ws)

    def name(self, name, ref):
        self.wb.defined_names[name] = DefinedName(name, attr_text=ref)

    def save(self, path):
        self.wb.active = 0
        self.wb.save(path)


def q(title):
    """A sheet name quoted for use inside a formula."""
    return "'" + title.replace("'", "''") + "'"


class Sheet:
    def __init__(self, book, ws):
        self.book = book
        self.ws = ws
        self.L = book.L
        self.align = "right" if self.L.he else "left"

    # ---- low level -------------------------------------------------------
    def _font(self, bold=False, color=INK, size=11, italic=False):
        return Font(name=FONT, bold=bold, color=color, size=size, italic=italic)

    def put(self, ref, value, fmt=None, bold=False, color=INK, fill=None, size=11,
            align=None, wrap=False, border=True, italic=False):
        c = self.ws[ref]
        c.value = value
        c.font = self._font(bold, color, size, italic)
        if fill:
            c.fill = PatternFill("solid", fgColor=fill)
        if fmt:
            c.number_format = fmt
        h = align or (None if isinstance(value, (int, float)) or (isinstance(value, str) and value.startswith("=")) else self.align)
        c.alignment = Alignment(horizontal=h, vertical="center", wrap_text=wrap, readingOrder=2 if self.L.he else 1)
        if border:
            c.border = BOX
        return c

    # ---- semantic cells --------------------------------------------------
    def title(self, text, sub, span):
        ws = self.ws
        last = col(span)
        ws.merge_cells(f"A1:{last}1")
        ws.merge_cells(f"A2:{last}2")
        self.put("A1", text, bold=True, color="FFFFFF", fill=GREEN, size=18, border=False)
        for i in range(2, span + 1):
            ws.cell(row=1, column=i).fill = PatternFill("solid", fgColor=GREEN)
        ws.row_dimensions[1].height = 34
        self.put("A2", sub, color=GREY, size=10, wrap=True, border=False)
        ws.row_dimensions[2].height = 30

    def section(self, ref, text, span=None):
        r = self.ws[ref].row
        c0 = self.ws[ref].column
        if span:
            self.ws.merge_cells(start_row=r, start_column=c0, end_row=r, end_column=c0 + span - 1)
        self.put(ref, text, bold=True, color=GREEN, size=12, border=False)

    def header(self, row, labels, start=1):
        for i, text in enumerate(labels):
            self.put(f"{col(start + i)}{row}", text, bold=True, color="FFFFFF", fill=GREEN, align="center", wrap=True)
        self.ws.row_dimensions[row].height = 32

    def label(self, ref, text, bold=False, fill=None, color=INK):
        return self.put(ref, text, bold=bold, fill=fill, color=color)

    def inp(self, ref, value=None, fmt=None, align=None):
        return self.put(ref, value, fmt=fmt, color=INPUT_FONT, fill=INPUT, align=align)

    def calc(self, ref, formula, fmt=None, bold=False, fill=None, color=INK, align=None):
        return self.put(ref, formula, fmt=fmt, bold=bold, fill=fill, color=color, align=align)

    def total(self, ref, formula, fmt=None):
        return self.put(ref, formula, fmt=fmt, bold=True, fill=GREEN_LIGHT)

    def note(self, ref, text, span=None):
        r = self.ws[ref].row
        c0 = self.ws[ref].column
        if span:
            self.ws.merge_cells(start_row=r, start_column=c0, end_row=r, end_column=c0 + span - 1)
        self.put(ref, text, color=GREY, size=9, wrap=True, border=False, italic=True)
        self.ws.row_dimensions[r].height = 28

    def kpi(self, label_ref, value_ref, text, formula, fmt=None):
        self.put(label_ref, text, bold=True, fill=GREEN_LIGHT)
        self.put(value_ref, formula, fmt=fmt, bold=True, fill=GREEN_LIGHT, size=12)

    # ---- table helpers ---------------------------------------------------
    def rows_input(self, first, last, columns, fmt_by_col=None, examples=None, align_by_col=None):
        """Yellow input cells for columns (letters) on rows first..last, with
        example values filled into the first rows."""
        fmt_by_col = fmt_by_col or {}
        align_by_col = align_by_col or {}
        for r in range(first, last + 1):
            ex = examples[r - first] if examples and r - first < len(examples) else None
            for i, c in enumerate(columns):
                v = ex[i] if ex is not None and i < len(ex) else None
                self.inp(f"{c}{r}", v, fmt=fmt_by_col.get(c), align=align_by_col.get(c))

    def rows_calc(self, first, last, letter, template, fmt=None, align=None):
        """template uses {r} for the row number."""
        for r in range(first, last + 1):
            self.calc(f"{letter}{r}", template.format(r=r), fmt=fmt, align=align)

    def dropdown(self, rng, source, allow_blank=True, prompt=None):
        dv = DataValidation(type="list", formula1=source, allow_blank=allow_blank, showDropDown=False)
        dv.error = self.L("יש לבחור מהרשימה (אפשר לערוך אותה בגיליון הרשימות)", "Pick from the list (you can edit it on the Lists sheet)")
        dv.errorTitle = self.L("ערך לא ברשימה", "Not in the list")
        dv.errorStyle = "warning"
        if prompt:
            dv.prompt = prompt
            dv.showInputMessage = True
        self.ws.add_data_validation(dv)
        dv.add(rng)

    def number_rule(self, rng, minimum=0):
        dv = DataValidation(type="decimal", operator="greaterThanOrEqual", formula1=str(minimum), allow_blank=True)
        dv.error = self.L("יש להזין מספר", "Enter a number")
        dv.errorStyle = "warning"
        self.ws.add_data_validation(dv)
        dv.add(rng)

    def status_colors(self, rng, mapping):
        """mapping: list of (text, fill, font)."""
        top_left = rng.split(":")[0]
        for text, fill, font in mapping:
            self.ws.conditional_formatting.add(rng, FormulaRule(
                formula=[f'{top_left}="{text}"'], fill=PatternFill("solid", fgColor=fill, bgColor=fill),
                font=Font(color=font, bold=True)))

    def rule(self, rng, formula, fill=RED_FILL, font=RED_FONT):
        self.ws.conditional_formatting.add(rng, FormulaRule(
            formula=[formula], fill=PatternFill("solid", fgColor=fill, bgColor=fill), font=Font(color=font)))

    def negative_red(self, rng):
        self.ws.conditional_formatting.add(rng, CellIsRule(
            operator="lessThan", formula=["0"], fill=PatternFill("solid", fgColor=RED_FILL, bgColor=RED_FILL), font=Font(color=RED_FONT, bold=True)))

    def databar(self, rng, color="5B9B87"):
        self.ws.conditional_formatting.add(rng, DataBarRule(start_type="num", start_value=0, end_type="num", end_value=1, color=color))

    def freeze(self, ref):
        self.ws.freeze_panes = ref


def lists_sheet(book, columns, title=None):
    """An editable 'Lists' sheet. columns: list of (name, heading, values, rows).
    Defines a workbook name for each list (whole reserved range, so items
    added later appear in the dropdowns)."""
    L = book.L
    s = book.sheet(title or L("רשימות", "Lists"), [26] * len(columns), landscape=False)
    s.title(L("רשימות לבחירה", "Dropdown lists"),
            L("אפשר לשנות, להוסיף או למחוק פריטים — הרשימות הנפתחות בגיליון מתעדכנות לבד.",
              "Change, add or remove items — the dropdowns in the workbook update by themselves."), max(len(columns), 2))
    s.ws.row_dimensions[2].height = 44
    refs = {}
    for i, (name, heading, values, rows) in enumerate(columns, start=1):
        c = col(i)
        s.header(4, [heading], start=i)
        for r in range(5, 5 + rows):
            v = values[r - 5] if r - 5 < len(values) else None
            s.inp(f"{c}{r}", v)
        ref = f"{q(s.ws.title)}!${c}$5:${c}${4 + rows}"
        book.name(name, ref)
        refs[name] = ref
    return s, refs
