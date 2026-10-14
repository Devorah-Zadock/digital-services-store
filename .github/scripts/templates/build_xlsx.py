#!/usr/bin/env python3
"""Builds DeskKit's downloadable Excel templates, in Hebrew (downloads/) and
English (downloads/en/), then:
  * recalculates every workbook with LibreOffice and stores the results as
    cached values, so phone/email previewers (which don't calculate) show
    real numbers instead of empty cells;
  * fails if any formula evaluates to an error (#REF!, #VALUE!, ...);
  * writes the sheet previews used by preview.html (js/product-preview-data.js
    is updated in place for the xlsx-* entries) and the catalog thumbnails
    (images/previews/xlsx-*.webp).

  python3 .github/scripts/templates/build_xlsx.py            # everything
  python3 .github/scripts/templates/build_xlsx.py --no-thumbs
  python3 .github/scripts/templates/build_xlsx.py --thumbs-only

Needs openpyxl, lxml, LibreOffice (soffice); thumbnails need Playwright.
"""
import datetime as dt
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
import zipfile

sys.path.insert(0, os.path.dirname(__file__))
import openpyxl  # noqa: E402
from lxml import etree  # noqa: E402

from xlsx_lib import Book, Lang  # noqa: E402
import xlsx_finance as F  # noqa: E402
import xlsx_ops as O  # noqa: E402
import xlsx_realestate as RE  # noqa: E402

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "..", ".."))

# slug (as in js/products-data.js) → (file name, builder)
TEMPLATES = [
    ("xlsx-income-expenses", "income-expenses.xlsx", F.income_expenses),
    ("xlsx-budget", "monthly-budget.xlsx", F.monthly_budget),
    ("xlsx-invoice", "invoice-template.xlsx", F.invoice),
    ("xlsx-expense", "expense-tracker.xlsx", F.expense_tracker),
    ("xlsx-clients", "client-tracker.xlsx", O.client_tracker),
    ("xlsx-cashflow", "cash-flow-forecast.xlsx", F.cash_flow),
    ("xlsx-inventory", "inventory-tracker.xlsx", O.inventory),
    ("xlsx-timeline", "project-timeline.xlsx", O.project_timeline),
    ("xlsx-pnl", "profit-loss-statement.xlsx", F.profit_loss),
    ("xlsx-timesheet", "timesheet-tracker.xlsx", O.timesheet),
    ("xlsx-breakeven", "breakeven-pricing.xlsx", F.breakeven),
    ("xlsx-salestargets", "sales-targets-tracker.xlsx", O.sales_targets),
    ("xlsx-vendors", "vendor-comparison.xlsx", O.vendor_comparison),
    ("xlsx-receivables", "receivables-tracker.xlsx", F.receivables),
    ("xlsx-event-budget", "event-budget.xlsx", O.event_budget),
    ("xlsx-savings", "savings-goals.xlsx", O.savings_goals),
    ("xlsx-shifts", "weekly-shift-schedule.xlsx", O.shift_schedule),
    ("xlsx-freelance-rate", "freelancer-hourly-rate.xlsx", F.freelancer_rate),
    ("xlsx-loan", "loan-calculator.xlsx", F.loan),
    ("xlsx-apartment-sales", "apartment-sales-tracker.xlsx", RE.apartment_sales),
    ("xlsx-construction-budget", "construction-budget.xlsx", RE.construction_budget),
]

NS = {"m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
ERRORS = ("#REF!", "#VALUE!", "#NAME?", "#DIV/0!", "#N/A", "#NUM!", "#NULL!", "Err:")


def recalc(src, workdir):
    """LibreOffice opens the workbook, recalculates and saves a copy."""
    out = os.path.join(workdir, "calc")
    os.makedirs(out, exist_ok=True)
    profile = os.path.join(workdir, "lo-profile")
    subprocess.run(["soffice", f"-env:UserInstallation=file://{profile}", "--headless", "--calc",
                    "--convert-to", "xlsx:Calc MS Excel 2007 XML", "--outdir", out, src],
                   check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=240)
    return os.path.join(out, os.path.basename(src))


def cached_values(calc_path):
    """{sheet title: {cell ref: value}} for every formula cell."""
    wf = openpyxl.load_workbook(calc_path, data_only=False)
    wv = openpyxl.load_workbook(calc_path, data_only=True)
    out = {}
    for ws in wf.worksheets:
        vals = {}
        wsv = wv[ws.title]
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("="):
                    vals[c.coordinate] = wsv[c.coordinate].value
        out[ws.title] = vals
    return out


def check_errors(slug, lang, values):
    bad = []
    for sheet, vals in values.items():
        for ref, v in vals.items():
            if isinstance(v, str) and v.startswith(ERRORS):
                bad.append(f"{sheet}!{ref}={v}")
    if bad:
        raise SystemExit(f"{slug} [{lang}]: formula errors: {bad[:12]}")


def inject(path, values):
    """Write the LibreOffice results into our own file as cached <v> values
    (the file keeps openpyxl's exact formatting; Excel still recalculates on
    open because fullCalcOnLoad is set)."""
    wb = openpyxl.load_workbook(path)
    titles = [ws.title for ws in wb.worksheets]
    tmp = path + ".tmp"
    with zipfile.ZipFile(path) as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        # workbook.xml order == sheetN order for openpyxl-written files
        for item in zin.infolist():
            data = zin.read(item.filename)
            m = re.match(r"xl/worksheets/sheet(\d+)\.xml$", item.filename)
            if m:
                title = titles[int(m.group(1)) - 1]
                data = inject_sheet(data, values.get(title, {}))
            zout.writestr(item, data)
    os.replace(tmp, path)


def inject_sheet(xml, vals):
    root = etree.fromstring(xml)
    for c in root.iter("{%s}c" % NS["m"]):
        f = c.find("m:f", NS)
        if f is None:
            continue
        ref = c.get("r")
        v = vals.get(ref)
        old = c.find("m:v", NS)
        if old is not None:
            c.remove(old)
        if v is None or v == "":
            c.set("t", "str")
            etree.SubElement(c, "{%s}v" % NS["m"]).text = ""
            continue
        if isinstance(v, bool):
            c.set("t", "b"); text = "1" if v else "0"
        elif isinstance(v, (int, float)):
            c.attrib.pop("t", None); text = repr(float(v)) if isinstance(v, float) else str(v)
        elif isinstance(v, (dt.datetime, dt.date)):
            c.attrib.pop("t", None)
            base = dt.datetime(1899, 12, 30)
            d = v if isinstance(v, dt.datetime) else dt.datetime(v.year, v.month, v.day)
            text = repr((d - base).total_seconds() / 86400)
        elif isinstance(v, dt.time):
            c.attrib.pop("t", None); text = repr((v.hour * 3600 + v.minute * 60 + v.second) / 86400)
        else:
            c.set("t", "str"); text = str(v)
        etree.SubElement(c, "{%s}v" % NS["m"]).text = text
    return etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)


# ------------------------------------------------------------------ previews
def fmt_value(v, number_format, lang):
    if v is None or v == "":
        return ""
    if isinstance(v, (dt.datetime, dt.date)):
        if "yyyy" in number_format and "dd" not in number_format:
            return v.strftime("%m/%Y")
        if number_format.strip() == "dd/mm":
            return v.strftime("%d/%m")
        return v.strftime("%d/%m/%Y")
    if isinstance(v, dt.time):
        return v.strftime("%H:%M")
    if isinstance(v, bool):
        return str(v)
    if isinstance(v, (int, float)):
        nf = number_format or ""
        if "%" in nf:
            return f"{v * 100:.1f}%" if "0.0%" in nf else f"{v * 100:.0f}%"
        if nf.startswith("hh"):
            secs = round(v * 86400)
            return f"{secs // 3600:02d}:{secs % 3600 // 60:02d}"
        dec = 2 if "0.00" in nf else (1 if "0.0" in nf and "0.00" not in nf else 0)
        s = f"{v:,.{dec}f}"
        if "₪" in nf:
            s = f"{s} ₪" if lang == "he" else f"₪{s}"
        return s
    return str(v)


def preview_rows(path, lang, sheet_title, max_rows, max_col=None):
    wb = openpyxl.load_workbook(path)
    wv = openpyxl.load_workbook(path, data_only=True)
    ws, wsv = wb[sheet_title], wv[sheet_title]
    ncol = max_col or min(ws.max_column, 14)
    spans, covered = {}, set()
    for mr in ws.merged_cells.ranges:
        if mr.min_row == mr.max_row:
            width = min(mr.max_col, ncol) - mr.min_col + 1
            spans[(mr.min_row, mr.min_col)] = width
            for cc in range(mr.min_col + 1, mr.max_col + 1):
                covered.add((mr.min_row, cc))
    rows = []
    for r in range(1, min(ws.max_row, max_rows) + 1):
        row = []
        for c in range(1, ncol + 1):
            if (r, c) in covered:
                continue
            cell = ws.cell(row=r, column=c)
            v = wsv.cell(row=r, column=c).value
            text = fmt_value(v, cell.number_format, lang)
            fill = cell.fill.fgColor.rgb if cell.fill and cell.fill.fill_type == "solid" else None
            fill = fill[-6:] if isinstance(fill, str) else None
            color = cell.font.color.rgb if cell.font and cell.font.color is not None and isinstance(cell.font.color.rgb, str) else None
            d = {"v": text}
            if fill and fill != "FFFFFF":
                d["fill"] = fill
            if cell.font.bold:
                d["bold"] = True
            if color and color[-6:] not in ("222222", "000000"):
                d["color"] = color[-6:]
            if cell.font.sz and float(cell.font.sz) != 11:
                d["size"] = float(cell.font.sz)
            if spans.get((r, c), 1) > 1:
                d["span"] = spans[(r, c)]
            row.append(d)
        # trim trailing empty cells, drop fully empty rows to a single []
        while row and not row[-1]["v"] and not row[-1].get("fill"):
            row.pop()
        rows.append(row)
    # collapse runs of empty rows at the end
    while rows and not rows[-1]:
        rows.pop()
    return rows


def update_preview_js(previews):
    p = os.path.join(ROOT, "js", "product-preview-data.js")
    src = open(p, encoding="utf-8").read()
    start = src.index("const PRODUCT_PREVIEW_DATA = {")
    end = src.rindex("};")
    data = json.loads(src[start + len("const PRODUCT_PREVIEW_DATA = "):end + 1])
    for k in [k for k in data if k.startswith("xlsx-")]:
        del data[k]
    data.update(previews)
    # decks keep the readable layout; sheets are one compact line each
    body = "{\n" + ",\n".join(
        f"  {json.dumps(k)}: " + (json.dumps(v, ensure_ascii=False, separators=(",", ":")) if k.startswith("xlsx-")
                                  else json.dumps(v, ensure_ascii=False, indent=2).replace("\n", "\n  "))
        for k, v in data.items()) + "\n}"
    open(p, "w", encoding="utf-8").write(src[:start] + "const PRODUCT_PREVIEW_DATA = " + body + ";" + src[end + 2:])


def thumbs_only():
    """Redraw the thumbnails from the preview data already in the JS file."""
    src = open(os.path.join(ROOT, "js", "product-preview-data.js"), encoding="utf-8").read()
    start = src.index("const PRODUCT_PREVIEW_DATA = {") + len("const PRODUCT_PREVIEW_DATA = ")
    data = json.loads(src[start:src.rindex("};") + 1])
    meta = {}
    for slug, fname, build in TEMPLATES:
        info = build(Book(Lang("he")))
        meta[slug] = {"header_row": info["rows"][0]}
    import thumbs as T
    T.render(meta, data)


def main():
    if "--thumbs-only" in sys.argv:
        return thumbs_only()
    thumbs = "--no-thumbs" not in sys.argv
    only = [a for a in sys.argv[1:] if not a.startswith("--")]
    previews = {}
    meta = {}
    with tempfile.TemporaryDirectory() as work:
        for lang in ("he", "en"):
            outdir = os.path.join(ROOT, "downloads", "" if lang == "he" else "en")
            os.makedirs(outdir, exist_ok=True)
            for slug, fname, build in TEMPLATES:
                if only and slug not in only:
                    continue
                L = Lang(lang)
                book = Book(L)
                info = build(book)
                path = os.path.join(work, lang + "-" + fname)
                book.save(path)
                vals = cached_values(recalc(path, work))
                check_errors(slug, lang, vals)
                inject(path, vals)
                final = os.path.join(outdir, fname)
                shutil.copyfile(path, final)
                main_title = info["main"].ws.title
                r0, r1 = info["rows"]
                rows = preview_rows(final, lang, main_title, max(r1 + 6, 22), info.get("maxcol"))
                key = slug if lang == "he" else slug + "@en"
                previews[key] = {"type": "xlsx", "sheetName": main_title, "rows": rows}
                meta[key] = {"file": os.path.relpath(final, ROOT), "header_row": r0, "sheet": main_title}
                print(f"  {lang} {fname}: {sum(len(v) for v in vals.values())} formulas OK")
    if not only:
        update_preview_js(previews)
    else:
        print("(single templates: preview data not rewritten)")
    if thumbs:
        import thumbs as T
        T.render({k: v for k, v in meta.items() if not k.endswith("@en")}, previews)


if __name__ == "__main__":
    main()
